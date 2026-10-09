"""Validate uploaded documents and extract plain text (PDF, DOCX, TXT, Markdown)."""

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from app.core.errors import DomainValidationError

ALLOWED: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
}


@dataclass(frozen=True)
class ValidatedFile:
    filename: str  # sanitized original name, for display only
    ext: str
    content_type: str


def safe_filename(name: str) -> str:
    base = PurePath(name.replace("\\", "/")).name
    base = re.sub(r"[^\w.\- ()]+", "_", base).strip(" .")
    return base[:200] or "upload"


def validate_upload(filename: str, data: bytes, max_bytes: int) -> ValidatedFile:
    """Check extension, size and file signature. Never trust the client's content type."""
    name = safe_filename(filename)
    ext = PurePath(name).suffix.lower()
    if ext not in ALLOWED:
        raise DomainValidationError(
            f"Unsupported file type '{ext or 'none'}'", {"allowed": sorted(ALLOWED)}
        )
    if not data:
        raise DomainValidationError("File is empty")
    if len(data) > max_bytes:
        raise DomainValidationError(
            "File is too large", {"max_bytes": max_bytes, "size": len(data)}
        )
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise DomainValidationError("File content is not a PDF")
    if ext == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise DomainValidationError("File content is not a DOCX document")
        except zipfile.BadZipFile as exc:
            raise DomainValidationError("File content is not a DOCX document") from exc
    if ext in (".txt", ".md", ".markdown"):
        try:
            data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise DomainValidationError("Text file must be UTF-8") from exc
    return ValidatedFile(name, ext, ALLOWED[ext])


def extract_text(data: bytes, ext: str) -> str:
    if ext == ".pdf":
        raw = _pdf_text(data)
    elif ext == ".docx":
        raw = _docx_text(data)
    else:
        raw = data.decode("utf-8")
    return normalize_text(raw)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text(extraction_mode="layout"))
        except Exception:  # noqa: BLE001 - fall back for PDFs the layout mode can't handle
            pages.append(page.extract_text())
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text
        style = (para.style.name or "").lower() if para.style is not None else ""
        if "list" in style and text.strip():
            text = "• " + text
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


BULLET_RE = re.compile(r"^\s*([•●▪◦■\-*–]|\d+[.)])\s+")


def normalize_text(raw: str) -> str:
    """Trim layout whitespace and re-join wrapped bullet lines.

    Layout extraction keeps visual indentation: a bullet continuation line is indented
    deeper than its bullet. Those are merged so one bullet == one line.
    """
    out: list[str] = []
    in_bullet = False
    for line in raw.replace("\r\n", "\n").replace("\x0c", "\n").split("\n"):
        stripped = re.sub(r"[ \t ]+", " ", line).strip()
        if not stripped:
            in_bullet = False
            continue
        indent = len(line) - len(line.lstrip())
        if BULLET_RE.match(line):
            out.append("• " + BULLET_RE.sub("", line).strip())
            in_bullet = True
        elif in_bullet and indent >= 3 and out:
            out[-1] += " " + stripped
        else:
            out.append(stripped)
            in_bullet = False
    return "\n".join(out)
