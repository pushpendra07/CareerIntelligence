"""Synthetic test documents (fictional candidate; no real personal data)."""

import io

SAMPLE_CV = """Asha Verma
Adobe Certified Expert | Magento Tech Lead
asha.verma@example.com | +91 98765 43210 | Jaipur, Rajasthan, India | linkedin.com/in/asha-verma-example
PROFESSIONAL SUMMARY
Technical Lead with 11+ years building Magento 2 and Adobe Commerce stores for retail and B2B clients.
TECHNICAL SKILLS
Commerce: Magento 2, Adobe Commerce Cloud, Hyva, MSI
Backend: PHP, MySQL, Redis, RabbitMQ, REST, GraphQL
Tools: Git, JIRA, Docker
PROFESSIONAL EXPERIENCE
Technical Lead
Acme Commerce Pvt Ltd, Jaipur | Jan 2020 – Present
• Led a team of 6 developers delivering Adobe Commerce Cloud projects.
• Reduced checkout latency by 40% using Redis and Varnish tuning.
• Mentored developers and ran code reviews.
Senior PHP Developer
Beta Web Studio, Pune | Jun 2014 – Dec 2019
• Built Magento 2 extensions and REST integrations with SAP.
• Migrated 3 stores from Magento 1 to Magento 2.
PROJECTS
Retail Store, Dubai (Adobe Commerce Cloud)
• Built a GraphQL layer for a headless PWA storefront.
CERTIFICATIONS
• Adobe Certified Expert – Adobe Commerce Developer
EDUCATION
B.Tech in Computer Science – Example Institute of Technology, Jaipur | 2010 – 2014
"""

SAMPLE_CV_V2 = SAMPLE_CV.replace(
    "Tools: Git, JIRA, Docker", "Tools: Git, JIRA, Docker, AWS, Kubernetes"
).replace("Technical Lead with 11+ years", "Tech Lead and architect-minded engineer with 11+ years")


def make_docx(text: str) -> bytes:
    import docx

    document = docx.Document()
    for line in text.strip().split("\n"):
        if line.startswith("• "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pdf(lines: list[str]) -> bytes:
    """A minimal single-page PDF with Helvetica text, built by hand (no PDF writer needed)."""

    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    ops = ["BT", "/F1 11 Tf", "14 TL", "50 780 Td"]
    for line in lines:
        ops.append(f"({esc(line)}) Tj T*")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()
