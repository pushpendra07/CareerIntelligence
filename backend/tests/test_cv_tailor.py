"""CV check and tailoring for one job: ATS score, DOCX edits that keep the format, nothing saved."""

import base64
import io
import zipfile
from pathlib import Path

import docx
import pytest
from docx.shared import Pt
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.cv import ats, tailor_docx, text_docx
from app.models.cv import CV
from app.services.job_service import JobInput, ingest_job

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
REQUIRED = ["PHP", "Magento", "MySQL", "GraphQL"]
PREFERRED = ["Docker", "Kubernetes"]


def make_cv(skills_one_per_line: bool = False, skills_heading: str = "Technical Skills") -> bytes:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Asha Verma · CV"
    d.add_paragraph("Asha Verma")
    d.add_paragraph("Senior Software Developer | E-commerce")
    d.add_paragraph("asha.verma@example.com")
    d.add_heading("Summary", level=1)
    d.add_paragraph("12 years building online stores with Magento 2 and PHP.")
    d.add_heading(skills_heading, level=1)
    if skills_one_per_line:
        for s in ("PHP", "Magento 2", "MariaDB"):
            d.add_paragraph(s, style="List Bullet")
    else:
        for label, items in (("Languages: ", "PHP, JavaScript"), ("Databases: ", "MariaDB, Redis"),
                             ("Platforms: ", "Shopify, Magento 2")):
            p = d.add_paragraph()
            p.add_run(label).bold = True
            p.add_run(items).font.size = Pt(10.5)
    d.add_heading("Experience", level=1)
    d.add_paragraph("Acme Retail — Senior Developer, Jan 2014 – Present")
    d.add_paragraph("Built Docker based deployments.", style="List Bullet")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def score(text: str) -> ats.AtsReport:
    return ats.score_cv(text, "Senior PHP Developer", REQUIRED, PREFERRED, 8)


def test_score_counts_found_implied_related_and_missing() -> None:
    report = score(tailor_docx.docx_text(make_cv()))
    status = {k.skill: (k.status, k.via) for k in report.keywords}
    assert status == {"PHP": ("found", None), "Magento": ("implied", "Magento 2"),
                      "MySQL": ("related", "MariaDB"), "GraphQL": ("missing", None),
                      "Docker": ("found", None), "Kubernetes": ("related", "Docker")}
    # required (1 + 0.75) / 4 * 50 + preferred 1/2 * 20 + title 15 + experience 15
    assert report.score == 62
    assert report.title_words_missing == [] and report.experience_found == 12


def test_parts_the_job_does_not_state_are_left_out() -> None:
    full = ats.score_cv("PHP developer", "PHP Developer", ["PHP"], [], None)
    assert full.score == 100 and set(full.parts) == {"required", "title"}
    unknown_years = ats.score_cv("PHP developer", None, ["PHP"], [], 5)
    assert unknown_years.score == 77 and unknown_years.notes
    custom = ats.score_cv("Knows Hyvä themes", None, ["Hyvä themes", "Luma"], [], None)
    assert [k.status for k in custom.keywords] == ["found", "missing"]


def _layout(data: bytes) -> list[tuple[str | None, list[bytes]]]:
    """Each paragraph's style and its runs' formatting: everything but the words."""
    doc = docx.Document(io.BytesIO(data))
    return [(p.style.name if p.style else None, [tailor_docx._rpr(r) for r in p.runs])
            for p in tailor_docx.paragraphs(doc)]


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _resaved(data: bytes) -> bytes:
    out = io.BytesIO()
    docx.Document(io.BytesIO(data)).save(out)
    return out.getvalue()


def test_tailor_adds_confirmed_skills_and_keeps_the_format() -> None:
    original = make_cv()
    result = tailor_docx.tailor(original, add=["Magento", "GraphQL"],
                                job_skills=REQUIRED + PREFERRED)
    text = tailor_docx.docx_text(result.data)
    assert "Platforms: Magento 2, Magento, Shopify" in text  # job skills first, kind kept
    assert "Languages: PHP, GraphQL, JavaScript" in text
    assert "Databases: MariaDB, Redis" in text  # nothing for the job here: untouched
    assert result.added == ["Magento", "GraphQL"] and result.not_added == []
    assert result.reordered_lines == 1
    assert score(text).score == 78

    # Same paragraphs, styles and run formats; every other part of the file is unchanged.
    assert _layout(result.data) == _layout(original)
    before, after = _parts(_resaved(original)), _parts(result.data)
    assert before.keys() == after.keys()
    assert [n for n in before if before[n] != after[n]] == ["word/document.xml"]


def test_tailor_never_adds_skills_nobody_confirmed() -> None:
    result = tailor_docx.tailor(make_cv(), add=[], job_skills=REQUIRED + PREFERRED)
    text = tailor_docx.docx_text(result.data)
    assert "GraphQL" not in text and "Kubernetes" not in text and result.added == []


def test_one_skill_per_line_gets_a_copy_of_the_last_line() -> None:
    original = make_cv(skills_one_per_line=True)
    result = tailor_docx.tailor(original, add=["GraphQL", "PHP"], job_skills=REQUIRED)
    doc = docx.Document(io.BytesIO(result.data))
    bullets = [p.text for p in doc.paragraphs if p.style is not None
               and p.style.name == "List Bullet"]
    assert bullets[:4] == ["PHP", "Magento 2", "MariaDB", "GraphQL"]  # PHP was already there
    assert result.added == ["GraphQL"]


def test_no_skills_section_changes_nothing_and_says_so() -> None:
    result = tailor_docx.tailor(make_cv(skills_heading="Toolbox"), add=["GraphQL"],
                                job_skills=REQUIRED)
    assert result.added == [] and result.not_added == ["GraphQL"]
    assert "No Skills section" in result.notes[0]


def test_headline_gets_the_job_title_in_front() -> None:
    result = tailor_docx.tailor(make_cv(), add=[], job_skills=REQUIRED,
                                headline_title="Senior PHP Developer")
    assert result.headline == "Senior PHP Developer | Senior Software Developer | E-commerce"
    assert result.headline in tailor_docx.docx_text(result.data).split("\n")


@pytest.fixture
def job_id(db: Session) -> int:
    job = ingest_job(db, JobInput(title="Senior PHP Developer", company_name="Acme Retail",
                                  url="https://jobs.acme.example/42", origin="import",
                                  required_skills=REQUIRED, preferred_skills=PREFERRED,
                                  experience_min=8)).job
    db.commit()
    return job.id


def test_api_checks_and_tailors_without_saving_anything(
        client: TestClient, db: Session, job_id: int, tmp_path: Path) -> None:
    cv = make_cv()
    checked = client.post(f"/api/v1/jobs/{job_id}/cv-check",
                          files={"file": ("Asha CV.docx", cv, DOCX)})
    assert checked.status_code == 200, checked.text
    body = checked.json()
    assert body["can_edit"] is True and body["report"]["score"] == 62
    assert body["suggested"] == ["Magento"]

    built = client.post(f"/api/v1/jobs/{job_id}/cv-tailor",
                        files={"file": ("Asha CV.docx", cv, DOCX)},
                        data={"add": ["Magento", "GraphQL", "Not a job skill"],
                              "use_title": "true"})
    assert built.status_code == 200, built.text
    out = built.json()
    assert (out["before"]["score"], out["after"]["score"]) == (62, 78)
    assert out["added"] == ["Magento", "GraphQL"]  # only the job's own skills can be added
    assert out["filename"] == "Asha-CV-Acme-Retail.docx"
    text = tailor_docx.docx_text(base64.b64decode(out["content_base64"]))
    assert text.split("\n")[1].startswith("Senior PHP Developer | ")

    # Nothing kept: no CV rows, no files in storage.
    assert db.scalar(select(func.count()).select_from(CV)) == 0
    storage = tmp_path / "storage"
    assert not storage.exists() or not any(storage.rglob("*"))


PDF_TEXT = """Asha Verma
Senior Software Developer | E-commerce
asha.verma@example.com | Pune, India
PROFESSIONAL SUMMARY
Senior developer with 12 years building online stores for retailers in India and the UAE, with
strong custom module development and integrations.
TECHNICAL SKILLS
Languages: PHP, JavaScript, TypeScript, HTML, CSS, Sass, Less, Python, Bash, SQL, Liquid
Templates, XML
Databases: MariaDB, Redis
EXPERIENCE
Acme Retail, Pune | Jan 2014 – Present
• Built Docker based deployments for twelve storefronts and moved the team to Git based
releases.
• Led code reviews."""


def test_text_cv_becomes_a_plain_word_copy_with_its_own_lines() -> None:
    paras = text_docx.paragraphs(PDF_TEXT)
    assert paras[:3] == [("name", "Asha Verma"),
                         ("headline", "Senior Software Developer | E-commerce"),
                         ("text", "asha.verma@example.com | Pune, India")]
    assert ("heading", "TECHNICAL SKILLS") in paras
    # Lines the PDF wrapped are joined again; separate entries stay separate.
    assert ("text", "Languages: PHP, JavaScript, TypeScript, HTML, CSS, Sass, Less, Python, Bash, "
                    "SQL, Liquid Templates, XML") in paras
    assert ("text", "Databases: MariaDB, Redis") in paras
    assert ("text", "Acme Retail, Pune | Jan 2014 – Present") in paras
    assert [t for k, t in paras if k == "bullet"] == [
        "Built Docker based deployments for twelve storefronts and moved the team to Git based "
        "releases.", "Led code reviews."]
    doc = docx.Document(io.BytesIO(text_docx.text_to_docx(PDF_TEXT)))
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["Asha Verma"] == "Title" and styles["EXPERIENCE"] == "Heading 1"
    assert styles["Led code reviews."] == "List Bullet"


def test_api_pdf_and_text_cvs_are_tailored_as_a_word_copy(client: TestClient,
                                                         job_id: int) -> None:
    text_cv = PDF_TEXT.encode()
    checked = client.post(f"/api/v1/jobs/{job_id}/cv-check",
                          files={"file": ("cv.txt", text_cv, "text/plain")}).json()
    assert checked["can_edit"] is True and checked["builds_copy"] is True
    assert checked["format"] == "txt"
    built = client.post(f"/api/v1/jobs/{job_id}/cv-tailor",
                        files={"file": ("cv.txt", text_cv, "text/plain")},
                        data={"add": ["GraphQL"]})
    assert built.status_code == 200, built.text
    out = built.json()
    assert out["added"] == ["GraphQL"] and out["filename"].endswith(".docx")
    assert out["notes"][0].startswith("Built as a Word copy of your TXT")
    text = tailor_docx.docx_text(base64.b64decode(out["content_base64"]))
    assert "GraphQL" in text and "Led code reviews." in text
    assert out["after"]["score"] > out["before"]["score"]


def test_api_job_without_skills_is_explained(client: TestClient, db: Session) -> None:
    bare = ingest_job(db, JobInput(title="Developer", company_name="Plain Co",
                                   url="https://plain.example/1", origin="import")).job
    db.commit()
    res = client.post(f"/api/v1/jobs/{bare.id}/cv-check",
                      files={"file": ("cv.docx", make_cv(), DOCX)})
    assert res.status_code == 422 and "no skills" in res.json()["error"]["message"]
